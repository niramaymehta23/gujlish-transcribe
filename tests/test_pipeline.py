"""Regression tests for output organization, resume and web exports."""

import io
import json
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np

from asr.cli import _decode_chunk, build_parser, run
from asr.postprocess import merge_segments
from asr.segmenter import AudioChunk
from asr.storage import new_run, run_paths
from evaluate import error_rate
from tests.helpers import DeferredExecutor
from web.server import create_app, transcript_docx


class TestMergeAndRedecode(unittest.TestCase):
    def test_forced_cut_removes_only_a_short_duplicate(self):
        refrain = "જય શ્રી કૃષ્ણ જય શ્રી કૃષ્ણ"
        segments = [{"text": refrain, "overlap": 0.0}, {"text": refrain + " આગળ", "overlap": 0.4}]
        # The six repeated words are real chanting; at most a few words of a 0.4 s overlap may be dropped.
        self.assertEqual(merge_segments(segments), f"{refrain} જય શ્રી કૃષ્ણ આગળ")
        segments = [{"text": "આજે machine learning", "overlap": 0.0}, {"text": "learning વિશે", "overlap": 0.4}]
        self.assertEqual(merge_segments(segments), "આજે machine learning વિશે")

    def test_fixed_mode_loops_are_redecoded_in_halves(self):
        class Backend:
            def __init__(self):
                self.calls = []

            def transcribe(self, samples, start, end, mode):
                self.calls.append((start, end))
                text = "એક બે " * 40 if len(self.calls) == 1 else " ".join(f"word{i}" for i in range(30))
                return {"start": start, "end": end, "text": text}

        backend = Backend()
        chunk = AudioChunk(0.0, 27.5, np.zeros(int(27.5 * 16_000), dtype=np.float32), True, 0.0)
        segment = _decode_chunk(backend, chunk, "mixed")
        self.assertEqual(len(backend.calls), 3)
        self.assertEqual(segment["flags"], ["redecoded_in_halves"])


class TestWorkflow(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "interview.wav"
        subprocess.run(
            [
                "ffmpeg",
                "-v",
                "error",
                "-f",
                "lavfi",
                "-i",
                "sine=frequency=440:duration=35",
                "-ar",
                "16000",
                str(self.source),
            ],
            check=True,
        )

    def tearDown(self):
        self.temp.cleanup()

    def test_resume_outputs_and_overlap(self):
        output = self.root / "transcripts" / "run"
        # Silero (if installed) treats a sine tone as non-speech and would skip every chunk; fixed chunking keeps the
        # two chunks this test relies on regardless of which detector is installed.
        args = build_parser().parse_args(
            [str(self.source), "flex", "--device", "cpu", "--chunking", "fixed", "--output-dir", str(output)]
        )

        class Backend:
            def __init__(self):
                self.calls = 0

            def transcribe(self, samples, start, end, mode):
                self.calls += 1
                if self.calls == 2:
                    raise RuntimeError("Simulated interruption")
                return {"start": start, "end": end, "text": "આજે machine learning"}

        with patch("asr.cli.choose_device", return_value="cpu"), patch("asr.cli.load_backend", return_value=Backend()):
            with self.assertRaisesRegex(RuntimeError, "interruption"):
                run(args)
        checkpoint = output / ".interview.flex.partial.json"
        self.assertEqual(len(json.loads(checkpoint.read_text())["segments"]), 1)

        class Resumed:
            def transcribe(self, samples, start, end, mode):
                return {"start": start, "end": end, "text": "machine learning વિશે વાત કરીશું"}

        events = []
        with patch("asr.cli.choose_device", return_value="cpu"), patch("asr.cli.load_backend", return_value=Resumed()):
            txt, data, plot = run(args, events.append)
        self.assertEqual(txt.name, "interview.flex.txt")
        self.assertEqual(txt.read_text().strip(), "આજે machine learning વિશે વાત કરીશું")
        self.assertIn("00:00:00,000 --> ", txt.with_suffix(".srt").read_text(encoding="utf-8"))
        self.assertEqual(len(json.loads(data.read_text())["segments"]), 2)
        self.assertFalse(checkpoint.exists())
        self.assertIsNone(plot)
        self.assertEqual(events[-1]["stage"], "complete")

    def test_reset_discards_the_old_matching_checkpoint(self):
        transcripts = self.root / "transcripts"
        argv = [str(self.source), "flex", "--device", "cpu", "--chunking", "fixed"]

        class Failing:  # finishes one chunk, so a checkpoint is saved, then stops
            calls = 0

            def transcribe(self, samples, start, end, mode):
                Failing.calls += 1
                if Failing.calls == 2:
                    raise RuntimeError("Simulated interruption")
                return {"start": start, "end": end, "text": "જૂનું"}

        class Working:
            def transcribe(self, samples, start, end, mode):
                return {"start": start, "end": end, "text": "નવું"}

        with patch("asr.cli.TRANSCRIPTS", transcripts), patch("asr.storage.TRANSCRIPTS", transcripts):
            with (
                patch("asr.cli.choose_device", return_value="cpu"),
                patch("asr.cli.load_backend", return_value=Failing()),
            ):
                with self.assertRaisesRegex(RuntimeError, "interruption"):
                    run(build_parser().parse_args(argv))
            self.assertEqual(len(list(transcripts.glob("*/*/*/.*.partial.json"))), 1)
            with (
                patch("asr.cli.choose_device", return_value="cpu"),
                patch("asr.cli.load_backend", return_value=Working()),
            ):
                run(build_parser().parse_args(argv + ["--reset"]))
            self.assertEqual(list(transcripts.glob("*/*/*/.*.partial.json")), [])

    def test_upload_collision_validation_and_downloads(self):
        executor = DeferredExecutor()
        with (
            patch("asr.storage.RECORDINGS", self.root / "recordings"),
            patch("asr.storage.TRANSCRIPTS", self.root / "transcripts"),
        ):
            app = create_app(executor=executor)
            client = app.test_client()
            self.assertEqual(client.get("/").status_code, 200)
            self.assertEqual(client.post("/api/jobs", data={}).status_code, 400)
            payload = self.source.read_bytes()
            ids = []
            for _ in range(2):
                response = client.post("/api/jobs", data={"audio": (io.BytesIO(payload), "વાત.wav"), "model": "flex"})
                self.assertEqual(response.status_code, 202)
                ids.append(response.json["id"])
            jobs = app.extensions["asr_jobs"]
            self.assertEqual(jobs[ids[0]]["name"], "વાત.wav")
            self.assertEqual(jobs[ids[1]]["name"], "વાત-2.wav")
            response = client.post("/api/jobs", data={"audio": (io.BytesIO(b"bad audio"), "bad.mp3")})
            self.assertEqual(response.status_code, 400)
            self.assertFalse((self.root / "recordings" / "bad.mp3").exists())
            self.assertEqual(client.post("/api/jobs", headers={"Origin": "https://other.example"}).status_code, 403)
            job = jobs[ids[0]]
            txt = Path(job["run_dir"]) / "વાત.flex.txt"
            text = "આજે machine learning વિશે વાત કરીશું\n"
            txt.write_text(text, encoding="utf-8")
            job.update(status="complete", txt=str(txt))
            self.assertEqual(client.get(f"/api/jobs/{ids[0]}").json["text"], text)
            with client.get(f"/api/jobs/{ids[0]}/download/txt") as download:
                self.assertEqual(download.data.decode(), text)
            response = client.get(f"/api/jobs/{ids[0]}/download/docx")
            self.assertEqual(response.status_code, 200)
            from docx import Document

            document = Document(io.BytesIO(response.data))
            self.assertEqual(document.paragraphs[0].text, text.strip())
            self.assertIn("attachment", response.headers["Content-Disposition"])
            self.assertEqual(client.get(f"/api/jobs/{ids[0]}/download/pdf").status_code, 404)

    def test_docx_multiline_unicode(self):
        from docx import Document

        text = "આજે machine learning\nગુજરાતી શબ્દો"
        doc = Document(transcript_docx(text, "Interview"))
        self.assertEqual([p.text for p in doc.paragraphs], text.splitlines())

    def test_independent_runs_and_plot_paths(self):
        with (
            patch("asr.storage.TRANSCRIPTS", self.root / "transcripts"),
            patch("asr.storage.PLOTS", self.root / "plots"),
        ):
            first = new_run(self.source, "flex", "mixed")
            second = new_run(self.source, "flex", "mixed")
            self.assertNotEqual(first, second)
            self.assertTrue(first.is_dir() and second.is_dir())
            txt, data, plot, checkpoint = run_paths(self.source, "flex", first)
            self.assertEqual(txt.parent, first)
            self.assertEqual(data.parent, first)
            self.assertEqual(checkpoint.parent, first)
            self.assertEqual(plot.relative_to(self.root / "plots").parent, first.relative_to(self.root / "transcripts"))

    def test_saved_history_survives_restart_in_date_order(self):
        root = self.root / "transcripts"
        # Different recording names: ordering must follow run time rather than name.
        for name, date in [("z_old", "20261014-120000-000001"), ("a_new", "20261014-130000-000001")]:
            folder = root / name / "flex-mixed" / date
            folder.mkdir(parents=True)
            (folder / f"{name}.flex.txt").write_text("ગુજરાતી", encoding="utf-8")
            (folder / f"{name}.flex.json").write_text(
                json.dumps({"source": f"{name}.wav", "model": "flex", "mode": "mixed", "duration": 3.0}),
                encoding="utf-8",
            )
        with patch("asr.storage.TRANSCRIPTS", root):
            client = create_app(executor=DeferredExecutor()).test_client()
            history = client.get("/api/history").json
            self.assertEqual([entry["name"] for entry in history], ["a_new.wav", "z_old.wav"])
            self.assertEqual(client.get(f"/api/jobs/{history[0]['id']}").json["text"], "ગુજરાતી")

    def test_evaluation_rates(self):
        self.assertEqual(error_rate(["આજે", "learning"], ["આજે", "machine", "learning"]), 0.5)
        self.assertEqual(error_rate(list("abc"), list("adc")), 1 / 3)
        self.assertEqual(error_rate([], []), 0.0)


if __name__ == "__main__":
    unittest.main()
