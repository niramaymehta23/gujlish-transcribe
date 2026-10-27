"""Web endpoints for the translation and summary stages, and the evaluation normalizer."""

import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from evaluate import error_rate, normalize
from tests.helpers import DeferredExecutor
from web.server import create_app, markdown_docx


class TestStages(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.run_dir = self.root / "transcripts" / "talk" / "flex-mixed" / "20261014-120000-000001"
        self.run_dir.mkdir(parents=True)
        self.txt = self.run_dir / "talk.flex.txt"
        self.json = self.run_dir / "talk.flex.json"
        self.txt.write_text("આજે વાત", encoding="utf-8")
        self.json.write_text(
            json.dumps(
                {
                    "source": "talk.wav",
                    "model": "flex",
                    "mode": "mixed",
                    "duration": 5.0,
                    "segments": [{"start": 0, "end": 5, "text": "આજે વાત"}],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        self.executor = DeferredExecutor()
        patcher = patch("asr.storage.TRANSCRIPTS", self.root / "transcripts")
        patcher.start()
        self.addCleanup(patcher.stop)
        # Snapshot os.environ so any API keys a test sets are rolled back afterwards.
        env = patch.dict("os.environ", {}, clear=False)
        env.start()
        self.addCleanup(env.stop)
        self.app = create_app(executor=self.executor)
        self.client = self.app.test_client()
        self.job_id = self.client.get("/api/history").json[0]["id"]

    def tearDown(self):
        self.temp.cleanup()

    def post(self, kind, body=None):
        return self.client.post(f"/api/jobs/{self.job_id}/{kind}", json=body or {})

    def test_history_lists_only_transcripts(self):
        (self.run_dir / "talk.flex.en.json").write_text('{"kind": "translation", "segments": []}', encoding="utf-8")
        (self.run_dir / "talk.flex.summary.json").write_text('{"kind": "summary"}', encoding="utf-8")
        self.assertEqual(len(self.client.get("/api/history").json), 1)

    def test_capabilities_never_expose_the_key(self):
        with patch.dict("os.environ", {"ANTHROPIC_API_KEY": "sk-secret"}), patch("web.server.LocalLLM") as local:
            local.return_value.base_url = "http://localhost:11434/v1"
            local.return_value.models.return_value = ["llama3"]
            response = self.client.get("/api/capabilities")
        self.assertEqual(response.json["claude"], True)
        self.assertEqual(response.json["local_models"], ["llama3"])
        self.assertNotIn("sk-secret", response.get_data(as_text=True))

    def test_translate_validation(self):
        self.assertEqual(self.post("translate", {"llm": "gpt"}).status_code, 400)
        self.assertEqual(self.post("translate", {"llm": "local", "model": "--evil flag"}).status_code, 400)
        with patch.dict("os.environ", {}, clear=True):
            self.assertEqual(self.post("translate", {"llm": "claude"}).status_code, 400)
            self.assertEqual(self.post("translate", {"llm": "gemini"}).status_code, 400)
        self.assertEqual(self.client.post("/api/jobs/nope/translate", json={}).status_code, 404)
        with patch.dict("os.environ", {"GEMINI_API_KEY": "g"}):
            self.assertEqual(self.post("translate", {"llm": "gemini"}).status_code, 202)
        self.assertEqual(
            self.client.post(f"/api/jobs/{self.job_id}/translate", headers={"Origin": "https://x.example"}).status_code,
            403,
        )

    def test_translate_queues_one_subprocess_and_rejects_a_second(self):
        response = self.post("translate", {"llm": "local", "model": "qwen3:14b"})
        self.assertEqual(response.status_code, 202)
        function, (job_id, key, command) = self.executor.tasks[0]
        self.assertEqual(key, "translation")
        self.assertTrue(command[1].endswith("translate.py"))
        self.assertIn("--llm-model=qwen3:14b", command)  # a single argument, never a loose value
        self.assertNotIn("ANTHROPIC_API_KEY", " ".join(command))
        self.assertEqual(self.post("translate", {"llm": "local"}).status_code, 409)
        status = self.client.get(f"/api/jobs/{self.job_id}").json
        self.assertEqual(status["translation"]["status"], "queued")

    def test_summary_needs_a_translation_and_a_finished_transcript(self):
        self.assertEqual(self.post("summarize", {"llm": "local"}).status_code, 409)
        (self.run_dir / "talk.flex.en.txt").write_text("Today's talk", encoding="utf-8")
        (self.run_dir / "talk.flex.en.json").write_text('{"kind": "translation", "segments": []}', encoding="utf-8")
        with patch.dict("os.environ", {"ANTHROPIC_API_KEY": "k"}):
            response = self.post("summarize", {"llm": "claude"})
        self.assertEqual(response.status_code, 202)
        command = self.executor.tasks[0][1][2]
        self.assertTrue(command[1].endswith("summarize.py"))
        self.assertIn("claude", command)

    def test_existing_files_count_as_complete_and_download(self):
        status = self.client.get(f"/api/jobs/{self.job_id}").json
        self.assertEqual(status["translation"]["status"], "idle")
        self.assertEqual(self.client.get(f"/api/jobs/{self.job_id}/download/en-txt").status_code, 404)
        (self.run_dir / "talk.flex.en.txt").write_text("Today's talk", encoding="utf-8")
        (self.run_dir / "talk.flex.summary.md").write_text("# Title\n\n- **Bold** point\n\n> quote", encoding="utf-8")
        status = self.client.get(f"/api/jobs/{self.job_id}").json
        self.assertEqual(status["translation"]["status"], "complete")
        self.assertEqual(status["translation_text"], "Today's talk")
        self.assertIn("# Title", status["summary_markdown"])
        with self.client.get(f"/api/jobs/{self.job_id}/download/en-txt") as response:
            self.assertEqual(response.data.decode(), "Today's talk")
        for kind in ("en-docx", "summary-docx"):
            response = self.client.get(f"/api/jobs/{self.job_id}/download/{kind}")
            self.assertEqual(response.status_code, 200)
            Document(io.BytesIO(response.data))
        with self.client.get(f"/api/jobs/{self.job_id}/download/summary-md") as response:
            self.assertEqual(response.status_code, 200)

    def test_markdown_docx_structure(self):
        document = Document(
            markdown_docx("# Title\n\n## Part\n\nPlain *emphasis* text\n\n- item one\n- item **two**", "x")
        )
        texts = [p.text for p in document.paragraphs]
        self.assertIn("Title", texts)
        self.assertIn("item one", texts)
        self.assertIn("item two", texts)
        self.assertNotIn("**", " ".join(texts))


class TestNormalize(unittest.TestCase):
    def test_normalize_removes_noise_but_keeps_gujarati_marks(self):
        self.assertEqual(normalize("  Hello,   World!  "), "hello world")
        self.assertEqual(normalize("આજે, વાત કરીશું।"), "આજે વાત કરીશું")
        self.assertEqual(normalize("ક્ષ"), normalize("ક્ષ"))
        self.assertEqual(normalize("é"), normalize("é"))  # composed and decomposed forms are equal

    def test_normalization_changes_the_score_only_for_punctuation_and_case(self):
        reference, prediction = "આજે, Machine learning.", "આજે machine learning"
        raw = error_rate(reference.split(), prediction.split())
        clean = error_rate(normalize(reference).split(), normalize(prediction).split())
        self.assertGreater(raw, 0)
        self.assertEqual(clean, 0)


if __name__ == "__main__":
    unittest.main()
